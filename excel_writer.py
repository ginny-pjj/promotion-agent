from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from models import AdvanceRecord, DoujiaRecord, Issue

DATE_FORMAT = "yyyy-mm-dd"


def _collect_issue_accounts(
    issues: list[Issue],
    doujia_records: list[DoujiaRecord],
    advance_records: list[AdvanceRecord],
) -> set[str]:
    accounts = {account for issue in issues for account in issue.accounts}
    if accounts:
        return accounts
    for issue in issues:
        if issue.source == "advance":
            accounts.update(item.account for item in advance_records)
        elif issue.source == "doujia":
            accounts.update(item.account for item in doujia_records)
    return accounts


def load_template_accounts(template: Path) -> set[str]:
    workbook = load_workbook(template, read_only=True, data_only=True)
    sheet = workbook.active
    header_row = find_header_row(sheet)
    headers = {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(sheet[header_row])
        if cell.value is not None
    }
    nickname_col = headers["抖音昵称"]
    accounts: set[str] = set()
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        account = str(row[nickname_col] or "").strip()
        if account:
            accounts.add(account)
    workbook.close()
    return accounts


def find_header_row(sheet) -> int:
    for row in sheet.iter_rows():
        values = {str(cell.value).strip() for cell in row if cell.value is not None}
        if {"抖音昵称", "抖加"}.issubset(values):
            return row[0].row
    raise ValueError("未找到包含“抖音昵称”和“抖加”的表头行")


def update_workbook(
    template: Path,
    output: Path,
    doujia_records: list[DoujiaRecord],
    advance_records: list[AdvanceRecord],
    issues: list[Issue],
) -> None:
    if template.resolve() == output.resolve():
        raise ValueError("输出文件不能覆盖原始模板")
    workbook = load_workbook(template)
    sheet = workbook.active
    header_row = find_header_row(sheet)
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[header_row]
        if cell.value is not None
    }
    required = {"抖音昵称", "抖加", "抖加支付人", "打款人", "打款日期"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Excel 缺少必要列: {', '.join(sorted(missing))}")

    doujia_by_account = {item.account: item for item in doujia_records}
    advance_by_account = {item.account: item for item in advance_records}
    issue_accounts = _collect_issue_accounts(issues, doujia_records, advance_records)
    expected_date_col = headers.get("预计打款日期")

    for row in sheet.iter_rows(min_row=header_row + 1):
        account = str(row[headers["抖音昵称"] - 1].value or "").strip()
        if not account:
            continue
        doujia = doujia_by_account.get(account)
        advance = advance_by_account.get(account)
        needs_review = account in issue_accounts
        remark_parts: list[str] = []

        if doujia:
            row[headers["抖加"] - 1].value = doujia.doujia_amount
            row[headers["抖加支付人"] - 1].value = doujia.doujia_payer
        if advance:
            row[headers["打款人"] - 1].value = advance.payer
            if advance.expected_payment_date:
                if expected_date_col:
                    _write_date_cell(row[expected_date_col - 1], advance.expected_payment_date)
                else:
                    remark_parts.append(f"预计打款：{advance.expected_payment_date.isoformat()}")
            needs_review = True

        if needs_review and "备注" in headers:
            remark_parts.insert(0, "待人工确认")
            row[headers["备注"] - 1].value = "；".join(remark_parts)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _write_date_cell(cell, value: date) -> None:
    cell.value = value
    cell.number_format = DATE_FORMAT
